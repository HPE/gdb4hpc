#!/usr/bin/python

import os
import sys
import csv
import select
import time
import openpyxl
import datetime
import smtplib
import email.message
import email.mime.text
from email.mime.multipart import MIMEMultipart
from email.mime.text import MIMEText

dt_tm=''

class Email_reports:


    def __init__(self):
        print "Creating Test Log & Emailing"

    def convertFormat(self, path):

      infile = os.path.join(path,'aesop.rep')
      outfile = os.path.join(path,'gdb4hpc_nightly_build_report.csv')
      
      with open(infile, 'r') as infile, \
        open(outfile, 'w') as outfile:
          reader = csv.reader(infile, delimiter='\t', skipinitialspace=True)
          writer = csv.writer(outfile, delimiter=' ',)
          for row in reader:
            writer.writerow(row)      
      
          infile.close()
          outfile.close()

    def update_Nigthly_Test_Log(self,path):

        # Process the input file
        infile = os.path.join(path,'gdb4hpc_nightly_build_report.csv')
        infile = open(infile, 'r')
        reader = csv.reader(infile, delimiter=' ', skipinitialspace=True)
        arInfile = []
        for row in reader:
          arInfile.append(row)
        

        # Prepair the outfile
        workbook = openpyxl.load_workbook(path + "/gdb4hpc_Nightly_Test_Log.xlsx")
        worksheets = workbook.get_sheet_names()

        for s in range(len(workbook.sheetnames)):
          if workbook.sheetnames[s] == 'template':
            break
        workbook.active = s
        
        # get active worksheet (template)
        templateWorkSheet = workbook.get_active_sheet()
 
        #get the "Number of rows
        template_row_count = templateWorkSheet.max_row
        #print template_row_count

        #get "Number of columns"
        template_column_count = templateWorkSheet.max_column
        #print template_column_count
        
        # copy the template to a new worksheet
        newWorkSheet = workbook.copy_worksheet(templateWorkSheet)
        
        # Activate new worksheet    
        workbook.active = newWorkSheet

        # name the new worksheet (date)
        currDateTime = str(datetime.datetime.now())
        dt = currDateTime.split(' ')[0]
        tm = currDateTime.split(' ')[1]
        hr = tm.split(':')[0]
        min = tm.split(':')[1]
        tmpsec = tm.split(':')[2]
        sec = tmpsec.split('.')[0]
        dt_tm = str(dt + "-" + hr + min + sec)
        newWorkSheet.title = dt_tm
        workSheetTitle = newWorkSheet.title
        newWorkSheetIndex = workbook.worksheets.index(workbook.get_sheet_by_name(newWorkSheet.title))
        #print workSheetTitle        
          
        workBookRow = 1
        arWorkBookColumn = ['A','B','C']
        errorCount = 0
        firstIt = 0
       
        #Update the workbook
        for row in arInfile:
           i = 0
           for element in row:
             location = arWorkBookColumn[i] + str(workBookRow)
             newWorkSheet[location] = element
             i = i + 1
             if i == 3:
               if element == 'N':
                 errorCount = errorCount + 1  
           workBookRow = workBookRow + 1
       
        #update worksheet with number of failures
        newWorkSheet['B50'] = errorCount

        workbook.save(path + "/gdb4hpc_Nightly_Test_Log.xlsx")

        #print newWorkSheetIndex
        return newWorkSheetIndex

    def emailReport(self,path,newWorkSheetIndex):
        workbook = openpyxl.load_workbook(path + "/gdb4hpc_Nightly_Test_Log.xlsx")
        workbook.active = newWorkSheetIndex
        newWorkSheet = workbook.get_active_sheet()

        BODY="""
        <meta http-equiv="Content-Type" content="text/html; charset=utf-8">
        <title>html title</title>
        <style type="text/css" media="screen">
        table {
            border-collapse: collapse;
            border-color: black;
            empty-cells:hide;
            border: 1px solid black;
            width: 800px;
        }
        td {
             border-color: black;
             empty-cells:hide;
             border: 1px solid black;
        }             
        </style>

        <table style=border: 1px black solid">
        <tbody>"""


        # Loop through the rows and columns, assemble the HTML string
        # containing report data
        r = newWorkSheet.max_row
        c = newWorkSheet.max_column

        for i in range(1,r+1):
          # create a new HTML row
          if i == 1:
            BODY = BODY + '<tr bgcolor=orange><B>'  
          elif 1 < i <= 10:
            BODY = BODY + '<tr bgcolor=beige><B>'  
          elif i == 11 or i >= r-1:
            BODY = BODY + '<tr bgcolor=yellow><B>'  
          else:
            BODY = BODY + '<tr bgcolor=white>'  
          for j in range(1,c+1):
              # create a new HTML cell
              if (newWorkSheet.cell(row=i,column=j).value) is None:
                BODY = BODY + '<td>' + '  ' + '</td>'
              else:
                if j == 2:
                  
                  if i == r and str(newWorkSheet.cell(row=i,column=j).value) > 0:
                    BODY = BODY + '<td align="center"><font color="red">' + str(newWorkSheet.cell(row=i,column=j).value) + '</font></td>'        
                  else:
                    BODY = BODY + '<td align="center">' + str(newWorkSheet.cell(row=i,column=j).value) + '</td>'        
                elif j == 3:
                  if str(newWorkSheet.cell(row=i,column=j).value) == 'N':
                    BODY = BODY + '<td align="center"><B><font color="red">' + str(newWorkSheet.cell(row=i,column=j).value) + '</font></td>'        
                  else:
                    BODY = BODY + '<td align="center"><B>' + str(newWorkSheet.cell(row=i,column=j).value) + '</td>'        
                else:
                  BODY = BODY + '<td>' + str(newWorkSheet.cell(row=i,column=j).value) + '</td>'        
          # end of row tag
          if i <= 11 or i >= r-1:
            BODY = BODY + '</B></tr>'
          else:
            BODY = BODY + '</tr>'
        # end the body & table
        BODY = BODY + "</tbody></table>"""  
         

        # email
        TO = 'pe_debuggers@cray.com'
        FROM = 'ulib@cray.com'
        msg = MIMEMultipart('alternative')
        msg['Subject'] = "GDB4HPC Nightly test results"
        msg['To'] = TO
        msg['From'] = FROM
        msg.preamble = """
        Your mail reader does not support the report format.
        Please visit us <a href="http://www.mysite.com">online</a>!"""


        HTML_BODY = MIMEText(BODY,'html')
        msg.attach(HTML_BODY)

        try:
          server = smtplib.SMTP('localhost')
        except Exception as e:
          print "exception:" +  str(e)
        #if __name__ == "__main__":
         # server.set_debuglevel(1)

        server.sendmail(FROM, [TO], msg.as_string())    
      
        workbook.save(path + "/gdb4hpc_Nightly_Test_Log.xlsx")

        return 0

def main(arg1):
    path = str(arg1)
    email = Email_reports()
    email.convertFormat(path)
    print "Updating the Test Log"
    newWorkSheetIndex = email.update_Nigthly_Test_Log(path)
    print "Emailing Test Summary"
    return_value = email.emailReport(path,newWorkSheetIndex)
    if return_value == 0:
      return 0
    else:
      return 1

if __name__ == "__main__":
  main(sys.argv[1:])

